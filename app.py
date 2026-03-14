import os
from dotenv import load_dotenv
import streamlit as st
from PyPDF2 import PdfReader
from docx import Document
import matplotlib.pyplot as plt
import re
from io import BytesIO
from fpdf import FPDF
from agent import ATSResumeAgent
from streamlit_option_menu import option_menu
import plotly.graph_objects as go

load_dotenv()

agent = ATSResumeAgent()

st.set_page_config(page_title="ResuParse AI ATS Analyzer", layout="wide", initial_sidebar_state="expanded")

# ----- Theme & Static Logo -----
st.markdown("""
    <style>
        /* Modern Base Font */
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
        html, body, [class*="css"] {
            font-family: 'Inter', sans-serif;
        }

        /* Glassmorphism Cards */
        .custom-card {
            background: rgba(255, 255, 255, 0.03);
            backdrop-filter: blur(10px);
            -webkit-backdrop-filter: blur(10px);
            border-radius: 16px;
            padding: 25px;
            margin-bottom: 25px;
            border: 1px solid rgba(255, 255, 255, 0.08);
            box-shadow: 0 4px 30px rgba(0, 0, 0, 0.1);
            transition: transform 0.3s ease, box-shadow 0.3s ease;
        }
        .custom-card:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 40px rgba(0, 191, 255, 0.15);
        }

        /* Modern Gradient Buttons */
        .stButton>button { 
            background: linear-gradient(135deg, #0066cc 0%, #00bfff 100%);
            color: white !important; 
            border-radius: 12px; 
            font-weight: 600; 
            padding: 0.5rem 1rem;
            border: none;
            box-shadow: 0 4px 15px rgba(0, 102, 204, 0.3);
            transition: all 0.3s ease; 
        }
        .stButton>button:hover { 
            transform: translateY(-2px) scale(1.02); 
            box-shadow: 0 6px 20px rgba(0, 191, 255, 0.4);
        }
        .stButton>button:active {
            transform: translateY(0) scale(0.98);
        }

        /* Upload area styling */
        [data-testid="stFileUploader"] {
            background: rgba(255, 255, 255, 0.02);
            border-radius: 12px;
            padding: 1rem;
            border: 1px dashed rgba(255, 255, 255, 0.2);
            transition: border-color 0.3s;
        }
        [data-testid="stFileUploader"]:hover {
            border-color: #00bfff;
        }

        /* Text Area overriding */
        .stTextArea textarea {
            border-radius: 10px;
            border: 1px solid rgba(255, 255, 255, 0.1);
            background-color: rgba(25, 28, 36, 0.7);
        }

        /* Layout spacing */
        .block-container {
            padding-top: 2rem;
            max-width: 1200px;
        }

        /* Typing Effect for Hero section */
        .typing-effect span {
            display: inline-block;
            animation: blink 1s infinite;
        }
        @keyframes blink {
            0% { opacity: 0; }
            50% { opacity: 1; }
            100% { opacity: 0; }
        }
        /* Landing Page Cards */
        .landing-card {
            background: rgba(255, 255, 255, 0.05);
            border-radius: 12px;
            padding: 20px;
            border: 1px solid rgba(255, 255, 255, 0.1);
            text-align: center;
            height: 100%;
        }
        .landing-card i {
            font-size: 24px;
            color: #00bfff;
            margin-bottom: 10px;
            display: block;
        }
    </style>
""", unsafe_allow_html=True)

st.markdown("""
    <div>
        <h1 style='color:#00bfff; margin-bottom: 0;'>ResuParse AI</h1>
        <p style='font-size:18px; margin-top: 0;'>Your AI guide to land your dream job <span class='typing-effect'>💡</span></p>
    </div>
""", unsafe_allow_html=True)

def extract_text_from_pdf(pdf_file):
    reader = PdfReader(pdf_file)
    return "".join(page.extract_text() or "" for page in reader.pages)

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    return "\n".join(para.text for para in doc.paragraphs)

def generate_pdf(feedback):
    if not isinstance(feedback, dict):
        # Fallback for old string-based feedback
        feedback = {"summary": str(feedback), "overall_match": 0, "category_scores": {}, "matched_keywords": [], "missing_keywords": [], "improvements": []}

    pdf = FPDF()
    pdf.add_page()
    
    # Header & Branding
    try:
        pdf.image("resuparseai.png", x=10, y=8, w=30)
    except:
        pass
    
    pdf.set_font("Arial", 'B', size=20)
    pdf.set_text_color(0, 191, 255) # Deep Sky Blue
    pdf.cell(0, 10, "ATS Resume Analysis Report", ln=True, align='C')
    pdf.ln(15)
    
    # 1. Executive Summary & Match Score
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", 'B', size=14)
    pdf.cell(0, 10, "1. Executive Summary", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    match_score = feedback.get("overall_match", 0)
    pdf.set_font("Arial", 'B', size=12)
    pdf.cell(40, 10, f"Overall Match Score: {match_score}%")
    
    # Simple Progress Bar in PDF
    pdf.set_fill_color(230, 230, 230)
    pdf.rect(60, pdf.get_y() + 2, 100, 6, 'F')
    bar_color = (75, 192, 192) if match_score > 75 else (255, 206, 86) if match_score > 50 else (255, 99, 71)
    pdf.set_fill_color(*bar_color)
    pdf.rect(60, pdf.get_y() + 2, match_score, 6, 'F')
    pdf.ln(10)
    
    pdf.set_font("Arial", size=11)
    summary = feedback.get("summary", "No summary available.")
    if isinstance(summary, list):
        summary = "\n".join(summary)
    
    # Sanitize for latin-1
    summary = summary.replace('\u2013', '-').replace('\u2014', '-').replace('\u2018', "'").replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
    summary = summary.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 8, summary)
    pdf.ln(10)
    
    # 2. Detailed Matching Matrix
    pdf.set_font("Arial", 'B', size=14)
    pdf.cell(0, 10, "2. Detailed Matching Matrix", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    cat_scores = feedback.get("category_scores", {})
    if cat_scores:
        pdf.set_font("Arial", size=11)
        for cat, score in cat_scores.items():
            cat_name = cat.replace('_', ' ').title()
            pdf.cell(60, 8, f"{cat_name}: {score}%", ln=True)
        pdf.ln(10)

    # 3. Skill Analysis
    pdf.set_font("Arial", 'B', size=14)
    pdf.cell(0, 10, "3. Skill Gap Analysis", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    pdf.set_font("Arial", 'B', size=12)
    pdf.set_text_color(46, 139, 87) # SeaGreen
    pdf.cell(0, 8, "Matched Skills:", ln=True)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", size=11)
    matched = feedback.get("matched_keywords", [])
    if matched:
        pdf.multi_cell(0, 8, ", ".join(matched))
    else:
        pdf.cell(0, 8, "No direct matches found.", ln=True)
    pdf.ln(5)
    
    pdf.set_font("Arial", 'B', size=12)
    pdf.set_text_color(178, 34, 34) # Firebrick
    pdf.cell(0, 8, "Missing Keywords:", ln=True)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", size=11)
    missing = feedback.get("missing_keywords", [])
    if missing:
        pdf.multi_cell(0, 8, ", ".join(missing))
    else:
        pdf.cell(0, 8, "No missing keywords identified!", ln=True)
    pdf.ln(10)

    # 4. Action Plan
    pdf.set_font("Arial", 'B', size=14)
    pdf.cell(0, 10, "4. Strategic Action Plan", ln=True)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(5)
    
    pdf.set_font("Arial", size=11)
    improvements = feedback.get("improvements", [])
    for idx, imp in enumerate(improvements):
        imp_clean = imp.replace('\u2013', '-').replace('\u2014', '-').replace('\u2018', "'").replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
        imp_clean = imp_clean.encode('latin-1', 'replace').decode('latin-1')
        pdf.multi_cell(0, 8, f"{idx+1}. {imp_clean}")
    
    return pdf.output(dest='S').encode('latin-1')

# -------- Session State Management --------
if "resume_file" not in st.session_state:
    st.session_state.resume_file = None
if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""
if "jd_text" not in st.session_state:
    st.session_state.jd_text = ""

# -------- Sidebar Navigation --------
with st.sidebar:
    st.image("resuparseai.png", width=280)
    page = option_menu(
        menu_title="Navigation",
        options=["📄 Resume Analyzer"],
        icons=["file-earmark-text"],
        menu_icon="cast",
        default_index=0,
        styles={
            "container": {"padding": "0!important", "background-color": "transparent"},
            "icon": {"color": "#00bfff", "font-size": "18px"}, 
            "nav-link": {"font-size": "15px", "text-align": "left", "margin":"5px", "--hover-color": "rgba(0, 191, 255, 0.1)"},
            "nav-link-selected": {"background-color": "rgba(0, 191, 255, 0.2)", "border-left": "4px solid #00bfff"},
        }
    )

# -------- Resume Analyzer Tab --------
if page == "📄 Resume Analyzer":
    # --- Landing Hero Section ---
    st.markdown("""
        <div class='custom-card' style='text-align: center; padding: 40px; background: linear-gradient(135deg, rgba(0, 102, 204, 0.1) 0%, rgba(0, 191, 255, 0.05) 100%);'>
            <h2 style='color: #00bfff; margin-bottom: 10px;'>Unlock Your Career Potential</h2>
            <p style='font-size: 1.1rem; opacity: 0.9;'>Optimize your resume for ATS algorithms and land more interviews with data-driven insights.</p>
        </div>
    """, unsafe_allow_html=True)

    # --- How it Works ---
    st.markdown("### 🛠️ How it Works")
    step1, step2, step3 = st.columns(3)
    with step1:
        st.markdown("<div class='landing-card'><strong>1. Target Job</strong><br><small>Paste the job description you're aiming for.</small></div>", unsafe_allow_html=True)
    with step2:
        st.markdown("<div class='landing-card'><strong>2. Upload Resume</strong><br><small>Upload or paste your current resume.</small></div>", unsafe_allow_html=True)
    with step3:
        st.markdown("<div class='landing-card'><strong>3. Get Insights</strong><br><small>Analyze and get an instant ATS-biased score.</small></div>", unsafe_allow_html=True)
    
    st.markdown("<br>", unsafe_allow_html=True)

    # --- Inputs Section ---
    st.markdown("### 📥 Input Your Details")
    col_jd, col_res = st.columns(2)
    
    with col_jd:
        st.markdown("#### 🎯 Job Description")
        jd_file = st.file_uploader("Upload JD", type=["pdf", "docx"], key="jd_file")
        jd_text = st.text_area("Or paste JD here", value=st.session_state.jd_text, height=200)
        if jd_text != st.session_state.jd_text:
            st.session_state.jd_text = jd_text
        jd = extract_text_from_pdf(jd_file) if jd_file and jd_file.name.lower().endswith("pdf") else extract_text_from_docx(jd_file) if jd_file else jd_text

    with col_res:
        st.markdown("#### 📄 Your Resume")
        rc_file = st.file_uploader("Upload Resume", type=["pdf", "docx"], key="resume_uploader")
        rc_text = st.text_area("Or paste Resume here", value=st.session_state.resume_text, key="resume_text_area", height=200)
        if rc_text != st.session_state.resume_text:
            st.session_state.resume_text = rc_text
        if rc_file is not None:
            st.session_state.resume_file = rc_file
        current_resume_file = st.session_state.resume_file
        resume_input = extract_text_from_pdf(current_resume_file) if current_resume_file and current_resume_file.name.lower().endswith("pdf") else extract_text_from_docx(current_resume_file) if current_resume_file else rc_text

    if resume_input and jd and st.button("🔍 Run Full Analysis"):
        with st.spinner("Analyzing your resume... This may take a moment."):
            feedback = agent.analyze_text(resume_input, jd)
            
            if not feedback:
                st.error("No analysis response received.")
            elif isinstance(feedback, dict):
                # Split feedback into 3 Tabs
                tab1, tab2, tab3 = st.tabs(["📊 Overview", "🎯 Skill Analysis", "🛠️ Action Plan"])
                
                with tab1:
                    # Overview Header
                    col_score, col_text = st.columns([1, 2])
                    
                    with col_score:
                        # Use direct 'overall_match' key for reliability
                        perc = feedback.get("overall_match", 0)
                        fig = go.Figure(go.Indicator(
                            mode = "gauge+number",
                            value = perc,
                            domain = {'x': [0, 1], 'y': [0, 1]},
                            title = {'text': "Match Score", 'font': {'color': 'white', 'size': 18}},
                            gauge = {
                                'axis': {'range': [None, 100], 'tickwidth': 1, 'tickcolor': "rgba(255,255,255,0.2)"},
                                'bar': {'color': "#00bfff"},
                                'bgcolor': "rgba(255,255,255,0.05)",
                                'borderwidth': 2,
                                'bordercolor': "rgba(255,255,255,0.1)",
                                'steps': [
                                    {'range': [0, 50], 'color': 'rgba(255, 99, 71, 0.2)'},
                                    {'range': [50, 75], 'color': 'rgba(255, 206, 86, 0.2)'},
                                    {'range': [75, 100], 'color': 'rgba(75, 192, 192, 0.2)'}
                                ]
                            }
                        ))
                        fig.update_layout(height=280, margin=dict(l=10, r=10, t=40, b=10), paper_bgcolor="rgba(0,0,0,0)", font={'color': "white"})
                        st.plotly_chart(fig, use_container_width=True)
                    
                    with col_text:
                        st.markdown("### 📝 AI Summary")
                        st.markdown(f"<div class='custom-card'>{feedback.get('summary', 'No summary available.')}</div>", unsafe_allow_html=True)
                    
                    st.markdown("---")
                    # Category Scores in columns
                    st.markdown("### 📋 Category Breakdown")
                    cat_scores = feedback.get("category_scores", {})
                    if cat_scores:
                        cols = st.columns(len(cat_scores))
                        for idx, (cat, score) in enumerate(cat_scores.items()):
                            with cols[idx]:
                                st.markdown(f"**{cat.replace('_', ' ').title()}**")
                                st.markdown(f"#### {score}%")
                                color = "#4bc0c0" if score > 75 else "#ffce56" if score > 50 else "#ff6347"
                                st.markdown(f"""
                                    <div style='background-color: rgba(255,255,255,0.1); border-radius: 5px; height: 10px; width: 100%;'>
                                        <div style='background-color: {color}; height: 100%; width: {score}%; border-radius: 5px;'></div>
                                    </div>
                                """, unsafe_allow_html=True)

                with tab2:
                    # Skills Analysis
                    col_list, col_chart = st.columns([1, 1.2])
                    
                    with col_list:
                        st.markdown("### 🎯 Skill Checklist")
                        matched = feedback.get("matched_keywords", [])
                        missing = feedback.get("missing_keywords", [])
                        
                        if matched:
                            with st.expander("✅ Matched Skills", expanded=True):
                                for k in matched:
                                    st.markdown(f"- ✨ {k}")
                        
                        if missing:
                            with st.expander("🔴 Missing Keywords", expanded=True):
                                for k in missing:
                                    st.markdown(f"- 🔴 {k}")
                        elif not matched and not missing:
                            st.info("Upload a resume and JD to see the skill analysis.")
                        else:
                            st.success("You have all the required keywords! perfection.")

                    with col_chart:
                        st.markdown("### 📊 Matching Insight")
                        m_count = len(feedback.get("matched_keywords", []))
                        miss_count = len(feedback.get("missing_keywords", []))
                        
                        if m_count + miss_count > 0:
                            fig = go.Figure(data=[
                                go.Bar(name='Found', x=['Skills'], y=[m_count], marker_color='#4bc0c0'),
                                go.Bar(name='Missing', x=['Skills'], y=[miss_count], marker_color='#ff6347')
                            ])
                            fig.update_layout(
                                barmode='group', 
                                height=350, 
                                paper_bgcolor="rgba(0,0,0,0)", 
                                plot_bgcolor="rgba(0,0,0,0)",
                                font={'color': "white"},
                                margin=dict(l=20, r=20, t=20, b=20)
                            )
                            st.plotly_chart(fig, use_container_width=True)
                        else:
                            st.caption("Complete an analysis to see the chart.")

                with tab3:
                    # Action Plan
                    st.markdown("### 🚀 Strategic Action Plan")
                    improvements = feedback.get("improvements", [])
                    if improvements:
                        for idx, imp in enumerate(improvements):
                            st.markdown(f"""
                                <div class='custom-card'>
                                    <strong>Step {idx+1}:</strong> {imp}
                                </div>
                            """, unsafe_allow_html=True)
                    
                    st.markdown("---")
                    st.markdown("### 🔧 Instant Bullet Point Optimizer")
                    st.markdown("Paste a weak bullet point below. We'll optimize it for the missing keywords identified above.")
                    
                    with st.container():
                        fix_text = st.text_area("Bullet point to improve:", height=100, placeholder="e.g. Responsible for managing a team and doing sales...")
                        if st.button("✨ Transform Bullet Point"):
                            if fix_text:
                                with st.spinner("Refining language..."):
                                    improved = agent.rephrase(fix_text)
                                    st.markdown("#### ✅ Optimized Result:")
                                    st.markdown(f"<div style='background-color: rgba(0, 191, 255, 0.1); padding: 15px; border-radius: 10px; border: 1px solid #00bfff;'>{improved}</div>", unsafe_allow_html=True)
                                    st.button("📋 Copy to clipboard", on_click=lambda: st.write("Text copied! (Demo only)"))
                            else:
                                st.warning("Please enter some text to rephrase.")

                st.markdown("---")
                st.download_button("📄 Download Full Feedback Report", data=generate_pdf(feedback), file_name="resume_analysis_report.pdf")
            else:
                st.markdown(feedback)

# -------- Magic Rephrase Tab --------
elif page == "✨ Magic Rephrase":
    st.header("🔮 Magic Rephrase")
    txt = st.text_area("Text to rephrase")
    if txt and st.button("Rephrase"):
        result = agent.rephrase(txt)
        st.success(result)

# -------- ATS Templates Tab --------
elif page == "📁 ATS Templates":
    st.header("📁 ATS Templates")
    st.markdown("Choose from these ATS-friendly resume templates:")
    templates = {
        "Modern Minimal": "https://docs.google.com/document/d/1NWFIz-EZ1ZztZSdXfrrcdffSzG-uermd/edit",
        "Elegant Blue": "https://docs.google.com/document/d/1xO7hvK-RQSb0mjXRn24ri3AiDrXx6qt8/edit",
        "Classic Chronological": "https://docs.google.com/document/d/1fAukvT0lWXns3VexbZjwXyCAZGw2YptO/edit"
    }
    for name, url in templates.items():
        st.markdown(f"**[{name}]({url})**")

# -------- Skill Gap Analyzer Tab --------
elif page == "📊 Skill Gap Analyzer":
    st.header("📊 Skill Gap Analyzer")
    st.markdown("### Upload Job Description or paste below")
    jd_file = st.file_uploader("Upload Job Description (PDF or DOCX)", type=["pdf", "docx"], key="jd_file2")
    jd_text = st.text_area("Or paste the Job Description here", value=st.session_state.jd_text, key="jd2")
    
    if jd_text != st.session_state.jd_text:
        st.session_state.jd_text = jd_text

    st.caption("📂 You can upload a job description file or paste it below.")
    jd = extract_text_from_pdf(jd_file) if jd_file and jd_file.name.lower().endswith("pdf") else extract_text_from_docx(jd_file) if jd_file else jd_text

    st.markdown("### Upload or paste your Resume")
    rf2_file = st.file_uploader("Upload resume", type=["pdf", "docx"], key="rf2")
    rf2_text = st.text_area("Or paste your Resume here", value=st.session_state.resume_text, key="rf2_text")
    if rf2_text != st.session_state.resume_text:
        st.session_state.resume_text = rf2_text

    if rf2_file:
        st.session_state.resume_file = rf2_file
        
    current_resume_file = st.session_state.resume_file
    resume_input = extract_text_from_pdf(current_resume_file) if current_resume_file and current_resume_file.name.lower().endswith("pdf") else extract_text_from_docx(current_resume_file) if current_resume_file else rf2_text

    if resume_input and jd and st.button("Analyze Skill Gaps"):
        with st.spinner("Analyzing skill gaps..."):
            result = agent.skill_gap_text(resume_input, jd)
            st.markdown("<div class='custom-card'>", unsafe_allow_html=True)
            st.markdown(result)
            st.markdown("</div>", unsafe_allow_html=True)

# -------- Cover Letter Gen Tab --------
elif page == "📝 Cover Letter Gen":
    st.header("📝 AI Cover Letter Generator")
    st.markdown("Generate a tailored cover letter based on your resume and target job.")
    
    jd_file = st.file_uploader("Upload Job Description (PDF or DOCX)", type=["pdf", "docx"], key="jd_cv_file")
    jd_text = st.text_area("Or paste the Job Description here", value=st.session_state.jd_text, key="jd_cv_text")
    
    if jd_text != st.session_state.jd_text:
        st.session_state.jd_text = jd_text

    jd = extract_text_from_pdf(jd_file) if jd_file and jd_file.name.lower().endswith("pdf") else extract_text_from_docx(jd_file) if jd_file else jd_text

    st.markdown("### Upload or paste your Resume")
    rc_cv_file = st.file_uploader("Upload Resume", type=["pdf", "docx"], key="rc_cv_file")
    rc_cv_text = st.text_area("Or paste your Resume here", value=st.session_state.resume_text, key="rc_cv_text")
    if rc_cv_text != st.session_state.resume_text:
        st.session_state.resume_text = rc_cv_text

    if rc_cv_file:
        st.session_state.resume_file = rc_cv_file

    current_resume_file = st.session_state.resume_file
    resume_input = extract_text_from_pdf(current_resume_file) if current_resume_file and current_resume_file.name.lower().endswith("pdf") else extract_text_from_docx(current_resume_file) if current_resume_file else rc_cv_text

    if resume_input and jd and st.button("Generate Cover Letter"):
        with st.spinner("Drafting your tailored cover letter..."):
            letter = agent.generate_cover_letter_text(resume_input, jd)
            st.markdown("<div class='custom-card'>", unsafe_allow_html=True)
            st.markdown(letter)
            st.markdown("</div>", unsafe_allow_html=True)
            st.download_button("📄 Download Cover Letter", data=generate_pdf(letter), file_name="cover_letter.pdf")

# -------- Interview Prep Tab --------
elif page == "🎤 Interview Prep":
    st.header("🎤 Interview Preparation Assistant")
    st.markdown("Generate customized interview questions based on your resume and the job description to help you prepare.")
    
    jd_file = st.file_uploader("Upload Job Description (PDF or DOCX)", type=["pdf", "docx"], key="jd_int_file")
    jd_text = st.text_area("Or paste the Job Description here", value=st.session_state.jd_text, key="jd_int_text")
    
    if jd_text != st.session_state.jd_text:
        st.session_state.jd_text = jd_text

    jd = extract_text_from_pdf(jd_file) if jd_file and jd_file.name.lower().endswith("pdf") else extract_text_from_docx(jd_file) if jd_file else jd_text

    st.markdown("### Upload or paste your Resume")
    rc_int_file = st.file_uploader("Upload Resume", type=["pdf", "docx"], key="rc_int_file")
    rc_int_text = st.text_area("Or paste your Resume here", value=st.session_state.resume_text, key="rc_int_text")
    if rc_int_text != st.session_state.resume_text:
        st.session_state.resume_text = rc_int_text
        
    if rc_int_file:
        st.session_state.resume_file = rc_int_file

    current_resume_file = st.session_state.resume_file
    resume_input = extract_text_from_pdf(current_resume_file) if current_resume_file and current_resume_file.name.lower().endswith("pdf") else extract_text_from_docx(current_resume_file) if current_resume_file else rc_int_text

    if resume_input and jd and st.button("Generate Interview Questions"):
        with st.spinner("Analyzing profile and generating questions..."):
            questions = agent.generate_interview_questions_text(resume_input, jd)
            st.markdown("<div class='custom-card'>", unsafe_allow_html=True)
            st.markdown(questions)
            st.markdown("</div>", unsafe_allow_html=True)
            st.download_button("📄 Download Interview Guide", data=generate_pdf(questions), file_name="interview_prep.pdf")
